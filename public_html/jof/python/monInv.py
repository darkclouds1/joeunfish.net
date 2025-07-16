import sys
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Border, Side, Alignment, Font
from openpyxl.utils import get_column_letter
import os

def format_date_with_weekday(date):
    if pd.isna(date):
        return "날짜 없음"
    weekdays = ['월', '화', '수', '목', '금', '토', '일']
    formatted_date = date.strftime("%m/%d") + '(' + weekdays[date.weekday()] + ')'
    return formatted_date

def process_excel(file_name):
    base_path = './data'
    file_path = os.path.join(base_path, file_name)

    if not os.path.exists(file_path):
        print(f"Error: File does not exist - {file_path}")
        return

    # 'openpyxl' 엔진을 사용하여 Excel 파일 읽기
    data = pd.read_excel(file_path, engine='openpyxl')
    data['출고일자'] = pd.to_datetime(data['출고일자'], errors='coerce')
    data['출고일자'] = data['출고일자'].apply(format_date_with_weekday)

    pivot_table = data.pivot_table(values='출고수량', index=['품번', '품명'], columns='출고일자', aggfunc='sum', fill_value=0)
    
    # 각 행의 합계를 계산하여 '합계' 열 추가
    pivot_table['합계'] = pivot_table.sum(axis=1)
    
    pivot_table.to_excel(file_path.replace('.xlsx', '_pivot.xlsx'), engine='openpyxl')

    wb = openpyxl.load_workbook(file_path.replace('.xlsx', '_pivot.xlsx'))
    ws = wb.active

    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    yellow_fill = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')

    for row in ws.iter_rows():
        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.font = Font(bold=False)
            if cell.column == 2:  # '품명' 열은 좌측 정렬
                cell.alignment = Alignment(horizontal='left', vertical='center')
            else:  # 나머지는 우측 정렬
                cell.alignment = Alignment(horizontal='right', vertical='center')

    for cell in ws[1]:
        cell.fill = yellow_fill
        cell.font = Font(bold=True)

    for col in ws.columns:
        max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col)
        adjusted_width = (max_length + 2)
        ws.column_dimensions[col[0].column_letter].width = adjusted_width

    ws.freeze_panes = 'A2'
    wb.save(file_path.replace('.xlsx', '_pivot.xlsx'))
    wb.close()
    print(f"File processed and saved as {file_path.replace('.xlsx', '_pivot.xlsx')}")

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python script.py <excel_file_name.xlsx>")
    else:
        file_name = sys.argv[1]
        process_excel(file_name)
